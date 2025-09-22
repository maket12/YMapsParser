import os

from dotenv import load_dotenv
from openpyxl import Workbook
from dataclasses import fields


def load_out_dir():
    load_dotenv()
    return os.getenv("OUT_DIR")


def export_to_excel(data: dict, filename: str = "output.xlsx"):
    """
    data — это словарь {org_id: Organization}, который возвращает parser.parse()
    """
    wb = Workbook()

    # --- Лист 1: Выдача ---
    ws1 = wb.active
    ws1.title = "Выдача"

    org_fields = fields(next(iter(data.values())))
    list1_fields = [f.name for f in org_fields if f.name not in ("reviews", "full_address", "lat", "lon", "phones", "site", "schedule", "is_open_now", "description", "photos_count", "has_video", "videos_count", "has_stories", "stories_titles", "services_count", "services_list", "services_prices", "services_has_photos", "services_with_photos", "services_has_covers", "services_description_filled", "has_news", "news_count", "news_titles", "news_texts", "news_dates", "has_features", "features_list", "card_url")]

    ws1.append(list1_fields)
    for org in data.values():
        row = [getattr(org, f) for f in list1_fields]
        ws1.append(row)

    # --- Лист 2: Организации ---
    ws2 = wb.create_sheet("Организации")

    list2_fields = [
        "org_id", "full_address", "lat", "lon", "phones", "site", "schedule",
        "is_open_now", "description", "photos_count", "has_video", "videos_count",
        "has_stories", "stories_titles", "services_count", "services_list",
        "services_prices", "services_has_photos", "services_with_photos",
        "services_has_covers", "services_description_filled", "has_news",
        "news_count", "news_titles", "news_texts", "news_dates", "has_features",
        "features_list", "card_url"
    ]

    ws2.append(list2_fields)
    for org in data.values():
        row = [getattr(org, f, None) for f in list2_fields]
        ws2.append(row)

    # --- Лист 3: Отзывы ---
    ws3 = wb.create_sheet("Отзывы")

    review_fields = [f.name for f in fields(next(iter(data.values())).reviews[0])] \
        if any(org.reviews for org in data.values()) else \
        ["review_id", "review_date", "stars", "text", "has_owner_reply", "owner_reply_text", "owner_reply_date"]

    ws3.append(["org_id"] + review_fields)
    for org in data.values():
        for review in org.reviews:
            row = [org.org_id] + [getattr(review, f) for f in review_fields]
            ws3.append(row)

    save_path = os.path.join(load_out_dir(), filename)
    wb.save(save_path)
    return save_path
